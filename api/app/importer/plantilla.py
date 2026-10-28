"""Plantilla descargable: las mismas hojas y encabezados que lee el importador.

Un solo origen para el contrato (`ENCABEZADOS` en `lectura.py`): si alguien cambia el
lector, la plantilla cambia con él y no aparecen dos formatos distintos.

La plantilla no trae datos de negocio (AGENTS.md: nada de datos inventados), solo los
encabezados y las instrucciones de la hoja `Leeme`.
"""

from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.importer.lectura import ENCABEZADOS

INSTRUCCIONES: tuple[tuple[str, str], ...] = (
    ("Caso", "Una fila por curso; el grado se deriva del dígito inicial del código."),
    ("Unidad", "Cada bloque es una franja con inicio y fin (HH:MM); la duración se deriva."),
    ("Jornada", "Lunes a viernes. La columna 'dia' escribe Lunes..Viernes."),
    ("Docentes", "max_sesiones_semana se cuenta en sesiones (bloques), no en horas."),
    ("Asignaturas docentes", "habilitado=1 registra la aptitud; 0 la omite."),
    ("Carga", "tipo_salon_requerido debe existir en la columna 'tipo' de Salones."),
    ("Salones", "Un nombre 'Aula 7A' marca el aula titular del curso 7A."),
    ("Disponibilidad", "1 disponible; 0 no disponible. Solo los 0 se guardan."),
    ("Cambio", "Escenario '¿y si...?': no altera la disponibilidad base."),
)

MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOMBRE_ARCHIVO = "houry-plantilla.xlsx"


def _escribir_encabezado(hoja: Worksheet, encabezado: tuple[str, ...]) -> None:
    hoja.append(list(encabezado))
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for posicion, titulo in enumerate(encabezado, start=1):
        hoja.column_dimensions[get_column_letter(posicion)].width = max(16, len(titulo) + 4)


def plantilla_bytes() -> bytes:
    """Genera la plantilla en memoria y devuelve sus bytes."""
    libro = openpyxl.Workbook()
    activa = libro.active
    assert activa is not None  # un Workbook nuevo siempre trae una hoja

    for posicion, (nombre, encabezado) in enumerate(ENCABEZADOS.items()):
        hoja = activa if posicion == 0 else libro.create_sheet()
        hoja.title = nombre
        _escribir_encabezado(hoja, encabezado)

    leeme = libro["Leeme"]
    for concepto, descripcion in INSTRUCCIONES:
        leeme.append([concepto, descripcion])
    leeme.column_dimensions["B"].width = 90

    flujo = BytesIO()
    libro.save(flujo)
    return flujo.getvalue()
