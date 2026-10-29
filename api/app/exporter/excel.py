"""Exportación a Excel (openpyxl, D-016). Una planilla por vista.

Grilla días × bloques con el texto de cada celda según la vista: asignatura y docente
por curso, curso y asignatura por docente o recurso. Sin hoja por bloque ni estilos
avanzados: es lo que una escuela imprime y pega en la pared.
"""

from __future__ import annotations

import uuid
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.validator.estado import EstadoHorario

DIAS = ("Lunes", "Martes", "Miércoles", "Jueves", "Viernes")

MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def celdas(
    sesiones: list[object], estado: EstadoHorario, por: str
) -> dict[tuple[int, uuid.UUID], str]:
    asignaturas = {asignatura.id: asignatura.nombre for asignatura in estado.asignaturas}
    docentes = {docente.id: docente.codigo for docente in estado.docentes}
    cursos = {curso.id: curso.codigo for curso in estado.cursos}
    asignaciones = {asignacion.id: asignacion for asignacion in estado.asignaciones}

    celdas: dict[tuple[int, uuid.UUID], str] = {}
    for sesion in sesiones:
        asignacion = asignaciones.get(sesion.asignacion_carga_id)
        asignatura = asignaturas.get(asignacion.asignatura_id) if asignacion else "?"
        if por == "docente" or por == "recurso":
            texto = f"{cursos.get(sesion.curso_id, '?')} · {asignatura}"
        else:
            texto = f"{asignatura} · {docentes.get(sesion.docente_id, '?')}"
        celdas[(sesion.dia, sesion.bloque_id)] = texto
    return celdas


def _hoja(
    hoja: Worksheet,
    titulo: str,
    mapa: dict[tuple[int, uuid.UUID], str],
    estado: EstadoHorario,
) -> None:
    hoja.title = titulo[:31]
    hoja.append(["Bloque", *DIAS])
    for bloque in sorted(estado.bloques, key=lambda bloque: bloque.orden):
        fila = [f"{bloque.codigo} ({bloque.inicio:%H:%M}-{bloque.fin:%H:%M})"]
        fila += [mapa.get((dia, bloque.id), "") for dia in range(1, 6)]
        hoja.append(fila)
    hoja.column_dimensions["A"].width = 22
    for columna in "BCDEF":
        hoja.column_dimensions[columna].width = 28


def objetivo(
    estado: EstadoHorario,
    sesiones: list[object],
    por: str,
    objetivo_id: uuid.UUID | None,
) -> list[object]:
    """Sesiones del alcance pedido. `general` lo resuelve cada exportador."""
    if por == "general" or objetivo_id is None:
        return list(sesiones)
    return [sesion for sesion in sesiones if getattr(sesion, f"{por}_id", None) == objetivo_id]


def exportar_xlsx(
    estado: EstadoHorario,
    sesiones: list[object],
    por: str,
    objetivo_id: uuid.UUID | None = None,
) -> bytes:
    """Devuelve el `.xlsx` en memoria.

    `por` es `curso`, `docente`, `recurso` o `general` (una hoja por curso).
    """
    libro = Workbook()
    libro.remove(libro.active)

    if por == "general":
        for curso in estado.cursos:
            propias = [sesion for sesion in sesiones if sesion.curso_id == curso.id]
            _hoja(libro.create_sheet(), curso.codigo, celdas(propias, estado, "curso"), estado)
    else:
        elegidas = objetivo(estado, sesiones, por, objetivo_id)
        titulo = {
            "curso": "Horario",
            "docente": "Docente",
            "recurso": "Aula",
        }.get(por, "Horario")
        _hoja(libro.create_sheet(), titulo, celdas(elegidas, estado, por), estado)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
