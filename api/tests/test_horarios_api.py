"""API de la fase 5: versiones, sesiones, validar, generar, deshacer y exportar."""

import uuid
from collections.abc import Iterator
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.db import get_session
from app.domain.carga import AsignacionCarga, DisponibilidadBloqueo
from app.domain.catalogo import Aptitud, Asignatura
from app.domain.enums import EstadoVersion, OrigenBloqueo, TipoBloque
from app.main import app
from app.repositories.carga import RepositorioAsignacionCarga, RepositorioDisponibilidad
from app.repositories.catalogo import (
    RepositorioAptitud,
    RepositorioAsignatura,
    RepositorioBloque,
)
from app.repositories.horario import RepositorioHorarioVersion, RepositorioSesion
from tests.factories import BaseHorario, crear_base_horario


@pytest.fixture
def cliente(sesion_db: Session) -> Iterator[TestClient]:
    app.dependency_overrides[get_session] = lambda: sesion_db
    try:
        with TestClient(app) as test_client:
            yield test_client
    finally:
        app.dependency_overrides.clear()


@pytest.fixture
def base(sesion_db: Session) -> BaseHorario:
    """Base mínima del horario + las aptitudes que el Excel real siempre trae."""
    armada = crear_base_horario(sesion_db)
    for docente in (armada.docente_a, armada.docente_b):
        RepositorioAptitud(sesion_db).agregar(
            Aptitud(docente_id=docente.id, asignatura_id=armada.asignatura.id)
        )
    return armada


def _url(base: BaseHorario, resto: str) -> str:
    return f"/api/horarios/{base.version.id}/{resto}"


def _colocar(cliente: TestClient, base: BaseHorario, **cambios: Any) -> Any:
    cuerpo = {
        "asignacion_carga_id": str(base.carga_a.id),
        "dia": 1,
        "bloque_id": str(base.bloque.id),
        **cambios,
    }
    return cliente.post(_url(base, "sesiones"), json=cuerpo)


def test_crear_version_numera_sobre_la_ultima(cliente: TestClient, base: BaseHorario) -> None:
    creada = cliente.post("/api/horarios", json={"nota": "prueba"})

    assert creada.status_code == 201
    assert creada.json()["numero"] == 2
    assert creada.json()["estado"] == "borrador"


def test_colocar_sesion_choca_por_curso(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    otra = RepositorioAsignatura(sesion_db).agregar(Asignatura(nombre="Inglés", color="#111111"))
    segunda = RepositorioAsignacionCarga(sesion_db).agregar(
        AsignacionCarga(
            curso_id=base.curso_a.id,
            asignatura_id=otra.id,
            docente_principal_id=base.docente_b.id,
            sesiones_semana=2,
            tipo_recurso_requerido=base.tipo_recurso.codigo,
        )
    )

    assert _colocar(cliente, base).status_code == 201
    choque = _colocar(cliente, base, asignacion_carga_id=str(segunda.id))

    assert choque.status_code == 409
    assert choque.json()["codigo"] == "CURSO_DOBLE"


def test_bloque_de_descanso_no_admite_sesiones(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    descanso = RepositorioBloque(sesion_db).agregar(
        replace(base.bloque, id=uuid.uuid4(), codigo="R1", tipo=TipoBloque.descanso)
    )

    respuesta = _colocar(cliente, base, bloque_id=str(descanso.id))

    assert respuesta.status_code == 400
    assert respuesta.json()["codigo"] == "BLOQUE_NO_CLASE"


def test_version_publicada_es_inmutable(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    publicada = RepositorioHorarioVersion(sesion_db).actualizar(
        replace(base.version, estado=EstadoVersion.publicado)
    )

    respuesta = _colocar(cliente, base)

    assert publicada.estado is EstadoVersion.publicado
    assert respuesta.status_code == 409
    assert respuesta.json()["codigo"] == "VERSION_PUBLICADA"


def test_validar_persiste_y_no_duplica(cliente: TestClient, base: BaseHorario) -> None:
    primera = cliente.post(_url(base, "validar"), json={}).json()["conflictos"]
    segunda = cliente.post(_url(base, "validar"), json={}).json()["conflictos"]

    # Una por curso: ninguno tiene sus 4 sesiones colocadas todavía.
    assert [item["codigo"] for item in primera] == ["MALLA_INCUMPLIDA", "MALLA_INCUMPLIDA"]
    assert len(primera) == len(segunda)
    assert cliente.get(_url(base, "conflictos")).json()["conflictos"] == segunda


def test_generar_persiste_el_plan(cliente: TestClient, base: BaseHorario) -> None:
    respuesta = cliente.post(_url(base, "generar"), json={"modo_docente": "habilitados"})

    assert respuesta.status_code == 200, respuesta.text
    cuerpo = respuesta.json()
    assert cuerpo["estado"] in {"OPTIMO", "FACTIBLE"}
    assert cuerpo["sesiones"] == 8
    assert cuerpo["conflictos"] == []
    assert len(cliente.get(f"/api/horarios/{base.version.id}/vista").json()["sesiones"]) == 8


def test_deshacer_y_rehacer(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    creada = _colocar(cliente, base).json()

    assert cliente.post(_url(base, "deshacer")).json()["deshecho"] is True
    assert cliente.get(f"/api/horarios/{base.version.id}/vista").json()["sesiones"] == []

    assert cliente.post(_url(base, "rehacer")).json()["rehecho"] is True
    sesiones = cliente.get(f"/api/horarios/{base.version.id}/vista").json()["sesiones"]
    assert [item["id"] for item in sesiones] == [creada["id"]]

    # Otra edición descarta la pila de rehacer.
    _colocar(cliente, base, dia=2)
    sin_redo = cliente.post(_url(base, "rehacer"))
    assert sin_redo.status_code == 409
    assert sin_redo.json()["codigo"] == "NADA_QUE_REHACER"


def test_publicar_bloquea_con_errores(cliente: TestClient, base: BaseHorario) -> None:
    respuesta = cliente.post(_url(base, "publicar"))

    assert respuesta.status_code == 409
    assert respuesta.json()["codigo"] == "CONFLICTOS_ABIERTOS"


def test_exportar_excel_abre_en_openpyxl(cliente: TestClient, base: BaseHorario) -> None:
    cliente.post(_url(base, "generar"), json={"modo_docente": "habilitados"})

    respuesta = cliente.get(_url(base, "exportar"), params={"por": "general"})

    assert respuesta.status_code == 200
    libro = openpyxl.load_workbook(BytesIO(respuesta.content))
    assert libro.sheetnames == ["6A", "6B"]
    assert libro["6A"]["A1"].value == "Bloque"
    assert libro["6A"]["B1"].value == "Lunes"


def test_arrastre_marca_franjas_no_disponibles(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    RepositorioDisponibilidad(sesion_db).agregar(
        DisponibilidadBloqueo(
            docente_id=base.docente_a.id,
            dia=1,
            bloque_id=base.bloque.id,
            origen=OrigenBloqueo.personal,
        )
    )

    celdas = cliente.get(
        _url(base, "arrastre"), params={"asignacion_id": str(base.carga_a.id)}
    ).json()["celdas"]

    assert len(celdas) == 5
    lunes = next(celda for celda in celdas if celda["dia"] == 1)
    assert lunes["estado"] == "bloqueada"
    assert lunes["motivo"]
    # Los otros días no chocan con nada: la malla incompleta no marca la celda.
    assert {celda["estado"] for celda in celdas if celda["dia"] != 1} == {"valida"}


def test_propuestas_salen_del_greedy(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    RepositorioDisponibilidad(sesion_db).agregar(
        DisponibilidadBloqueo(
            docente_id=base.docente_a.id,
            dia=1,
            bloque_id=base.bloque.id,
            origen=OrigenBloqueo.personal,
        )
    )
    _colocar(cliente, base)
    conflicto = cliente.post(_url(base, "validar"), json={}).json()["conflictos"][0]

    respuesta = cliente.post(f"/api/conflictos/{conflicto['id']}/propuestas")

    assert respuesta.status_code == 200, respuesta.text
    propuestas = respuesta.json()["propuestas"]
    assert propuestas and propuestas[0]["origen"] == "solver"
    assert propuestas[0]["cambios"][0]["tipo"] in {
        "mover_sesion",
        "cambiar_docente",
        "cambiar_recurso",
    }

    aceptada = cliente.post(f"/api/propuestas/{propuestas[0]['id']}/aceptar")
    assert aceptada.status_code == 200
    codigos = {
        item["codigo"]
        for item in cliente.get(_url(base, "conflictos")).json()["conflictos"]
    }
    assert "DOCENTE_NO_DISPONIBLE" not in codigos


def test_version_inexistente_devuelve_404(cliente: TestClient) -> None:
    respuesta = cliente.get("/api/horarios/00000000-0000-0000-0000-000000000000")

    assert respuesta.status_code == 404
    assert respuesta.json()["codigo"] == "NO_ENCONTRADO"


def test_escenario_duplicado_devuelve_409(cliente: TestClient) -> None:
    assert cliente.post("/api/escenarios", json={"nombre": "D2 sin martes"}).status_code == 201
    respuesta = cliente.post("/api/escenarios", json={"nombre": "D2 sin martes"})

    assert respuesta.status_code == 409
    assert respuesta.json()["codigo"] == "NOMBRE_DUPLICADO"


def test_escenario_activo_es_el_de_la_version(cliente: TestClient) -> None:
    escenario = cliente.post("/api/escenarios", json={"nombre": "Sin D1"}).json()

    version = cliente.post("/api/horarios", json={"escenario_id": escenario["id"]}).json()

    assert version["escenario_id"] == escenario["id"]


def test_flujo_completo_con_el_excel_real(cliente: TestClient) -> None:
    """Importar -> generar -> publicar -> exportar, el camino que muestra el pitch."""
    datos = Path(__file__).resolve().parents[2] / "data" / "Datos_prueba_Gnosoft.xlsx"
    media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    subida = cliente.post(
        "/api/importaciones", files={"archivo": (datos.name, datos.read_bytes(), media)}
    )
    assert subida.status_code == 200, subida.text
    confirmada = cliente.post(f"/api/importaciones/{subida.json()['id']}/confirmar")
    assert confirmada.status_code == 200, confirmada.text

    version = cliente.post("/api/horarios", json={"nota": "demo"}).json()
    generado = cliente.post(
        f"/api/horarios/{version['id']}/generar", json={"modo_docente": "habilitados"}
    )
    assert generado.status_code == 200, generado.text
    cuerpo = generado.json()
    assert cuerpo["sesiones"] == 120
    assert [item for item in cuerpo["conflictos"] if item["gravedad"] == "error"] == []

    publicada = cliente.post(f"/api/horarios/{version['id']}/publicar")
    assert publicada.status_code == 200, publicada.text
    assert publicada.json()["estado"] == "publicado"

    exportado = cliente.get(f"/api/horarios/{version['id']}/exportar?por=general")
    libro = openpyxl.load_workbook(BytesIO(exportado.content))
    assert len(libro.sheetnames) == 4


# --- pin de sesión (D-017) ----------------------------------------------------------


def test_pin_se_guarda_y_se_alterna(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    creada = _colocar(cliente, base, fijada=True).json()
    assert creada["fijada"] is True

    liberada = cliente.patch(
        _url(base, f"sesiones/{creada['id']}"), json={"fijada": False}
    )
    assert liberada.status_code == 200
    assert liberada.json()["fijada"] is False

    sesiones = cliente.get(f"/api/horarios/{base.version.id}/vista").json()["sesiones"]
    assert sesiones[0]["fijada"] is False


def test_generar_respeta_la_sesion_fijada(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    """Con un pin, el solver mueve el resto pero no la fijada."""
    fijada = _colocar(cliente, base, dia=5, fijada=True).json()
    _colocar(cliente, base, dia=1)

    respuesta = cliente.post(_url(base, "generar"), json={"modo_docente": "habilitados"})

    assert respuesta.status_code == 200, respuesta.text
    sesiones = cliente.get(f"/api/horarios/{base.version.id}/vista").json()["sesiones"]
    ancla = next(item for item in sesiones if item["id"] == fijada["id"])
    assert (ancla["dia"], ancla["bloque_id"]) == (5, str(base.bloque.id))
    assert ancla["fijada"] is True


def test_deshacer_restaura_el_pin(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    creada = _colocar(cliente, base).json()

    cliente.patch(_url(base, f"sesiones/{creada['id']}"), json={"fijada": True})
    cliente.post(_url(base, "deshacer"))

    sesiones = cliente.get(f"/api/horarios/{base.version.id}/vista").json()["sesiones"]
    assert sesiones[0]["fijada"] is False


def test_persistir_plan_sobrevive_a_un_intercambio_de_franjas(
    sesion_db: Session, base: BaseHorario
) -> None:
    """Regresión: A→B y B→A choca contra los UNIQUE si se actualiza fila por fila."""
    from dataclasses import replace as reemplazar

    from app.domain.horario import Sesion as SesionDominio
    from app.services.horario import persistir_plan
    from app.solver.resultado import EstadoResultado, Resultado

    segunda = RepositorioBloque(sesion_db).agregar(
        replace(base.bloque, id=uuid.uuid4(), codigo="B2", orden=2)
    )
    repo = RepositorioSesion(sesion_db)
    una = repo.agregar(
        SesionDominio(
            horario_version_id=base.version.id,
            asignacion_carga_id=base.carga_a.id,
            curso_id=base.curso_a.id,
            docente_id=base.docente_a.id,
            dia=1,
            bloque_id=base.bloque.id,
        )
    )
    otra = repo.agregar(
        SesionDominio(
            horario_version_id=base.version.id,
            asignacion_carga_id=base.carga_a.id,
            curso_id=base.curso_a.id,
            docente_id=base.docente_a.id,
            dia=2,
            bloque_id=segunda.id,
        )
    )

    intercambiadas = (
        reemplazar(una, dia=2, bloque_id=segunda.id),
        reemplazar(otra, dia=1, bloque_id=base.bloque.id),
    )
    persistir_plan(
        sesion_db,
        base.version,
        Resultado(estado=EstadoResultado.FACTIBLE, sesiones=intercambiadas),
    )

    guardadas = {
        item.id: (item.dia, item.bloque_id)
        for item in repo.listar_por_version(base.version.id)
    }
    assert guardadas[una.id] == (2, segunda.id)
    assert guardadas[otra.id] == (1, base.bloque.id)


def test_nueva_version_copia_una_publicada_sin_tocarla(
    cliente: TestClient, base: BaseHorario, sesion_db: Session
) -> None:
    """D-021: lo publicado no se edita. Se crea un borrador a partir de él."""
    creada = _colocar(cliente, base, fijada=True).json()
    RepositorioHorarioVersion(sesion_db).actualizar(
        replace(base.version, estado=EstadoVersion.publicado)
    )

    respuesta = cliente.post("/api/horarios", json={"copiar_de": str(base.version.id)})

    assert respuesta.status_code == 201, respuesta.text
    borrador = respuesta.json()
    assert borrador["numero"] == 2
    assert borrador["estado"] == "borrador"

    copiadas = RepositorioSesion(sesion_db).listar_por_version(uuid.UUID(borrador["id"]))
    assert len(copiadas) == 1
    assert copiadas[0].fijada is True
    assert copiadas[0].id != uuid.UUID(creada["id"])

    # La publicada sigue intacta y con su sesión.
    originales = RepositorioSesion(sesion_db).listar_por_version(base.version.id)
    assert [str(item.id) for item in originales] == [creada["id"]]


def test_copiar_una_version_inexistente_da_404(cliente: TestClient, base: BaseHorario) -> None:
    respuesta = cliente.post(
        "/api/horarios", json={"copiar_de": "00000000-0000-0000-0000-000000000000"}
    )

    assert respuesta.status_code == 404
    assert respuesta.json()["codigo"] == "NO_ENCONTRADO"
